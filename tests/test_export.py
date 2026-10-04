"""Excel files: the builders (app/services/export.py) and the file tools of the agent.

"Now" is Sunday 4 Oct 2026 10:00 Tehran = 12 Mehr 1405 (Mehr 1405 = 23 Sep – 22 Oct 2026).
"""

import json
from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from app.agent.actions import ActionLog
from app.agent.tools import ToolContext, build_registry
from app.config import Calendar
from app.db.models import Expense
from app.services.expenses import ExpenseService
from app.services.export import (
    TableSheet,
    build_table,
    date_cell,
    file_slug,
    range_label,
    sheet_name,
    table_cell,
)
from app.services.reminders import ReminderService, to_utc
from app.services.settings import (
    KEY_NIGHTLY,
    KEY_NIGHTLY_TIME,
    SettingsService,
)

TZ = ZoneInfo("Asia/Tehran")
NOW = datetime(2026, 10, 4, 10, 0, tzinfo=TZ)
REGISTRY = build_registry()


@pytest.fixture
async def ctx(config, sessionmaker):
    async with sessionmaker() as session:
        settings = SettingsService(session, Calendar.JALALI)
        expenses = ExpenseService(session, TZ, "toman")
        reminders = ReminderService(session, TZ, config.day_times)
        reminders.now = lambda: NOW  # type: ignore[method-assign]
        yield ToolContext(
            config=config,
            calendar=Calendar.JALALI,
            now=NOW,
            user_text="(test)",
            expenses=expenses,
            reminders=reminders,
            settings=settings,
            actions=ActionLog(session, expenses, reminders, settings),
        )


async def call(ctx, name: str, **args):
    return await REGISTRY.execute(name, json.dumps(args, ensure_ascii=False), ctx)


async def add(ctx, description, amount, when, category="Groceries", quantity=None, unit=None):
    found = await ctx.expenses.category_by_name(category)
    ctx.expenses.session.add(
        Expense(
            amount=amount,
            category_id=found.id,
            description=description,
            quantity=quantity,
            unit=unit,
            spent_at=to_utc(when),
            raw_text=description,
        )
    )
    await ctx.expenses.session.commit()


@pytest.fixture
async def spent(ctx):
    await add(ctx, "ماست", 200_000, datetime(2026, 10, 4, 9, 0, tzinfo=TZ))
    await add(ctx, "بنزین", 100_000, datetime(2026, 10, 3, 20, 0, tzinfo=TZ), "Fuel", 10, "L")
    await add(ctx, "رستوران", 1_500_000, datetime(2026, 9, 25, 13, 0, tzinfo=TZ), "Restaurant")
    await add(ctx, "نان", 50_000, datetime(2026, 9, 20, 8, 0, tzinfo=TZ))  # 29 Shahrivar
    await add(ctx, "کتاب", 300_000, datetime(2026, 8, 1, 18, 0, tzinfo=TZ), "Education")
    return ctx


def workbook(outcome):
    assert outcome.card is not None and outcome.card.kind == "file"
    return load_workbook(BytesIO(outcome.card.attachment))


def rows(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


# --- Labels and cells ---


def test_range_labels():
    assert range_label(date(2026, 9, 23), date(2026, 10, 23), Calendar.JALALI) == (
        "Mehr 1405",
        "1405-07",
    )
    assert range_label(date(2026, 10, 1), date(2026, 11, 1), Calendar.GREGORIAN) == (
        "October 2026",
        "2026-10",
    )
    assert range_label(date(2026, 3, 21), date(2027, 3, 21), Calendar.JALALI) == ("1405", "1405")
    assert range_label(date(2026, 10, 4), date(2026, 10, 5), Calendar.JALALI) == (
        "1405/07/12",
        "1405-07-12",
    )
    assert range_label(date(2026, 9, 23), date(2026, 9, 28), Calendar.JALALI) == (
        "1405/07/01 – 1405/07/05",
        "1405-07-01_1405-07-05",
    )


def test_dates_follow_the_calendar():
    assert date_cell(date(2026, 10, 4), Calendar.JALALI) == "1405/07/12"  # sortable text
    assert date_cell(date(2026, 10, 4), Calendar.GREGORIAN) == date(2026, 10, 4)  # a real date


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("120,000", 120000),
        ("۱۲۰۰۰۰", 120000),
        ("12.5", 12.5),
        (30, 30),
        ("15%", 0.15),
        ("09121234567", "09121234567"),  # a phone number stays text
        ("1234567890123456789", "1234567890123456789"),  # card numbers too
        ("Saturday", "Saturday"),
        (True, "Yes"),
        (None, ""),
    ],
)
def test_table_cells(value, expected):
    assert table_cell(value) == expected


def test_sheet_and_file_names():
    used: set[str] = set()
    assert sheet_name("Plan: week 1/2", used) == "Plan week 12"
    assert sheet_name("plan week 12", used) == "plan week 12 (2)"
    assert sheet_name("", used) == "Sheet"
    assert len(sheet_name("x" * 40, used)) == 31
    assert file_slug("برنامه ورزشی هفتگی") == "برنامه-ورزشی-هفتگی"
    assert file_slug("???") == "vira-table"


def test_table_formats_and_totals():
    file = build_table(
        "Prices",
        [
            TableSheet(
                "Shops",
                ["Item", "Price", "Year", "Off"],
                [["Milk", "120,000", "2026", "10%"], ["Bread", "۴۵۰۰۰", "2025", "5%"]],
                totals=True,
            )
        ],
    )
    ws = load_workbook(BytesIO(file.content)).active
    assert rows(ws) == [
        ["Item", "Price", "Year", "Off"],
        ["Milk", 120000, 2026, 0.1],
        ["Bread", 45000, 2025, 0.05],
        ["Total", 165000, None, None],  # years and percentages are not added up
    ]
    assert ws["B2"].number_format == "#,##0"
    assert ws["C2"].number_format == "0"  # no thousands separator for years
    assert ws["D2"].number_format == "0.0%"
    assert file.filename == "Prices.xlsx"


def test_text_stays_text():
    file = build_table("x", [TableSheet("s", ["Note"], [["=HYPERLINK(1)"], [chr(7) + "bell"]])])
    ws = load_workbook(BytesIO(file.content)).active
    assert ws["A2"].value == "=HYPERLINK(1)" and ws["A2"].data_type == "s"  # not a formula
    assert ws["A3"].value == "bell"  # control characters dropped


# --- export_expenses ---


async def test_this_month_by_default(spent):
    outcome = await call(spent, "export_expenses")
    assert outcome.result["count"] == 3 and outcome.result["total"] == 1_800_000
    assert outcome.result["range"] == {"from": "2026-09-23", "to": "2026-10-22"}
    assert outcome.card.data["filename"] == "vira-expenses-1405-07.xlsx"
    assert outcome.card.data["caption"]["title"] == "Expenses · Mehr 1405"
    assert "[done: sent Excel file vira-expenses-1405-07.xlsx" in outcome.note

    wb = workbook(outcome)
    assert wb.sheetnames == ["Expenses", "By category", "By day"]
    expenses = rows(wb["Expenses"])
    assert expenses[0] == [
        "Date", "Weekday", "Time", "Description", "Category", "Quantity", "Amount (toman)",
    ]  # fmt: skip
    assert expenses[1] == [
        "1405/07/03", "Friday", "13:00", "رستوران", "Restaurant", None, 1500000,
    ]  # fmt: skip
    assert expenses[2][3:6] == ["بنزین", "Fuel", "10 L"]
    assert expenses[-1][0] == "Total" and expenses[-1][-1] == 1_800_000
    assert wb["Expenses"]["G2"].number_format == "#,##0"

    categories = rows(wb["By category"])
    assert categories[1] == ["Restaurant", 1500000, pytest.approx(1_500_000 / 1_800_000), 1]
    assert categories[-1] == ["Total", 1800000, 1, 3]
    days = rows(wb["By day"])
    assert len(days) == 1 + 12 + 1  # 1–12 Mehr (up to today) and the total
    assert days[1] == ["1405/07/01", "Wednesday", 0, 0]
    assert len(wb["By category"]._charts) == 1 and len(wb["By day"]._charts) == 1


async def test_month_names_years_and_dates(spent):
    shahrivar = await call(spent, "export_expenses", month="شهریور")
    assert shahrivar.result["count"] == 1
    assert shahrivar.card.data["caption"]["title"] == "Expenses · Shahrivar 1405"

    october = await call(spent, "export_expenses", month="October")
    assert october.result["count"] == 2  # 3 and 4 October
    assert october.card.data["filename"] == "vira-expenses-2026-10.xlsx"
    assert october.card.data["caption"]["title"] == "Expenses · October 2026"

    year = await call(spent, "export_expenses", year="1405")
    assert year.result["count"] == 5
    assert workbook(year).sheetnames == ["Expenses", "By category", "By month"]
    months = rows(workbook(year)["By month"])
    assert [m[0] for m in months[1:-1]] == [
        "Farvardin 1405", "Ordibehesht 1405", "Khordad 1405", "Tir 1405", "Mordad 1405",
        "Shahrivar 1405", "Mehr 1405",
    ]  # fmt: skip

    days = await call(spent, "export_expenses", from_date="1405-07-01", to_date="1405-07-05")
    assert days.result["count"] == 1
    assert days.card.data["caption"]["title"] == "Expenses · 1405/07/01 – 1405/07/05"


async def test_filters_and_choices(spent):
    groceries = await call(spent, "export_expenses", categories=["groceries"])
    assert groceries.result["count"] == 1
    fuel = await call(spent, "export_expenses", query="بنزین", period="this_month")
    assert fuel.result["count"] == 1
    lean = await call(spent, "export_expenses", columns=["description", "amount"], summaries=[])
    wb = workbook(lean)
    assert wb.sheetnames == ["Expenses"]
    assert rows(wb["Expenses"])[0] == ["Description", "Amount (toman)"]

    unknown = await call(spent, "export_expenses", categories=["Pets"])
    assert "unknown category" in unknown.result["error"] and "Groceries" in unknown.result["hint"]
    bad = await call(spent, "export_expenses", summaries=["hourly"])
    assert "hourly" in bad.result["error"]


async def test_weeks_start_on_saturday_inside_the_range(spent):
    outcome = await call(spent, "export_expenses", summaries=["week"])
    weeks = rows(workbook(outcome)["By week"])
    assert [w[0] for w in weeks[1:-1]] == [
        "1405/07/01 – 1405/07/03",  # Mehr starts on a Wednesday
        "1405/07/04 – 1405/07/10",
        "1405/07/11 – 1405/07/17",
    ]
    assert [w[1:] for w in weeks[1:-1]] == [[1_500_000, 1], [0, 0], [300_000, 2]]


async def test_no_expenses_no_file(spent):
    outcome = await call(spent, "export_expenses", from_date="1405-06-01", to_date="1405-06-10")
    assert outcome.card is None and outcome.result["count"] == 0


# --- export_reminders ---


async def test_reminders_file(ctx):
    await call(
        ctx, "create_reminder", subject="دکتر", when_text="فردا ساعت ۲",
        start="2026-10-05T14:00", important=True,
    )  # fmt: skip
    await call(
        ctx, "create_reminder", subject="باشگاه", when_text="شنبه بعد ساعت ۸",
        start="2026-10-10T08:00", repeat="weekly",
    )  # fmt: skip
    upcoming = await call(ctx, "export_reminders")
    assert upcoming.result["count"] == 2
    assert upcoming.card.data["filename"] == "vira-reminders-upcoming.xlsx"
    table = rows(workbook(upcoming).active)
    assert table[0] == [
        "Date", "Weekday", "Time", "Subject", "Repeat", "Alerts", "Important", "Status",
    ]  # fmt: skip
    assert table[1][:4] == ["1405/07/13", "Monday", "14:00", "دکتر"]
    assert table[1][6:] == ["Yes", "Active"]
    assert table[2][4] == "Every Saturday"

    next_week = await call(ctx, "export_reminders", period="next_week")
    assert next_week.result["count"] == 1  # Sat 10 – Fri 16 Oct (Jalali week)
    nothing = await call(ctx, "export_reminders", query="dentist")
    assert nothing.card is None and nothing.result["count"] == 0


# --- make_spreadsheet ---


async def test_any_table(ctx):
    outcome = await call(
        ctx,
        "make_spreadsheet",
        title="برنامه ورزشی هفتگی",
        sheets=[
            {
                "name": "Plan",
                "columns": ["Day", "Workout", "Minutes"],
                "rows": [["Saturday", "Running", "30"], ["Monday", "Gym", 45]],
                "totals": True,
            }
        ],
    )
    assert outcome.card.data["filename"] == "برنامه-ورزشی-هفتگی.xlsx"
    assert rows(workbook(outcome)["Plan"]) == [
        ["Day", "Workout", "Minutes"],
        ["Saturday", "Running", 30],
        ["Monday", "Gym", 45],
        ["Total", None, 75],
    ]


async def test_table_rows_as_objects(ctx):
    outcome = await call(
        ctx,
        "make_spreadsheet",
        title="Shopping",
        sheets=[{"columns": [], "rows": [{"Item": "Milk", "Price": "120,000"}, {"Item": "Eggs"}]}],
    )
    assert rows(workbook(outcome).active) == [["Item", "Price"], ["Milk", 120000], ["Eggs", None]]


async def test_table_limits(ctx):
    empty = await call(ctx, "make_spreadsheet", title="x", sheets=[])
    assert "no sheets" in empty.result["error"]
    huge = await call(
        ctx, "make_spreadsheet", title="x", sheets=[{"columns": ["a"], "rows": [["1"]] * 2001}]
    )
    assert "too many rows" in huge.result["error"]


# --- Nightly report and briefing settings from the chat ---


async def test_nightly_settings_from_the_chat(ctx):
    outcome = await call(ctx, "update_settings", nightly_report_time="۱۱ شب")
    assert outcome.result["changed"] == {KEY_NIGHTLY_TIME: "23:00", KEY_NIGHTLY: "on"}
    assert (await ctx.settings.nightly_time(datetime.min.time())).hour == 23

    off = await call(ctx, "update_settings", nightly_report=False)
    assert off.result["changed"] == {KEY_NIGHTLY: "off"}
    assert not await ctx.settings.nightly_enabled()

    evening = await call(ctx, "update_settings", nightly_report_time="9:30")
    assert evening.result["changed"]["nightly_report_time"] == "21:30"  # the evening, of course
    morning = await call(ctx, "update_settings", nightly_report_time="9 am")
    assert "between 12:00 and 23:59" in morning.result["error"]
    briefing = await call(ctx, "update_settings", morning_briefing_time="7")
    assert briefing.result["changed"]["morning_briefing_time"] == "07:00"

    await ctx.actions.undo(off.card.action_id)  # ↩️ Undo brings the report back
    assert await ctx.settings.nightly_enabled()
